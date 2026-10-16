"""CSV и XLSX: колонки, форматы чисел, оформление, условное форматирование, графики."""

from __future__ import annotations

import csv
import zipfile
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from kns_report.aggregate import DailyReport, build_report
from kns_report.cli import required_tags, write_outputs
from kns_report.config import ReportConfig
from kns_report.csv_out import write_alarms_csv, write_hourly_csv
from kns_report.influx import parse_annotated_csv, rows_to_activations, rows_to_series
from kns_report.tagmodel import AlarmDef, TagInfo
from kns_report.timeutil import day_bounds
from kns_report.xlsx_out import FIRST_DATA_ROW, HEADER_ROW, SHEET_ALARMS, SHEET_CHARTS, SHEET_HOURLY, SHEET_SUMMARY

from .conftest import FIXTURES, GENERATED_AT, MSK

EXPECTED_COLUMNS = [
    "hour_start",
    "level_min_m",
    "level_avg_m",
    "level_max_m",
    "volume_m3",
    "p101_run_h",
    "p101_starts",
    "p102_run_h",
    "p102_starts",
    "p103_run_h",
    "p103_starts",
    "alarms_total",
    "alarms_high",
    "alarms_medium",
    "alarms_low",
    "data_coverage_pct",
]


@pytest.fixture()
def report(cfg: ReportConfig, tags: dict[str, TagInfo], alarms: list[AlarmDef]) -> DailyReport:
    rows = parse_annotated_csv((FIXTURES / "influx_day.csv").read_text(encoding="utf-8"))
    series = rows_to_series(cfg.influx, rows, required_tags(cfg.station, alarms, alarm_words=False))
    events = parse_annotated_csv((FIXTURES / "influx_day_alarms.csv").read_text(encoding="utf-8"))
    start, end = day_bounds(date(2026, 10, 15), MSK)
    return build_report(
        series,
        cfg.station,
        tags,
        alarms,
        start,
        end,
        "файл influx_day.csv",
        GENERATED_AT,
        end,
        rows_to_activations(cfg.influx, events),
    )


def test_hourly_csv(report: DailyReport, tmp_path: Path) -> None:
    path = tmp_path / "day.csv"
    write_hourly_csv(report, path)
    raw = path.read_bytes()
    assert not raw.startswith(b"\xef\xbb\xbf") and b"\r\n" not in raw
    rows = list(csv.DictReader(path.open(encoding="utf-8"), delimiter=";"))
    assert list(rows[0]) == EXPECTED_COLUMNS
    assert len(rows) == 24
    assert rows[0]["hour_start"] == "2026-10-15T00:00:00+03:00"
    assert rows[14]["alarms_total"] == str(report.hours[14].alarms_total)
    assert float(rows[8]["volume_m3"]) == pytest.approx(report.hours[8].volume_m3 or 0.0, abs=0.005)
    assert rows[3]["data_coverage_pct"] == str(round(report.hours[3].coverage * 100, 1))
    assert all(len(r["level_avg_m"].split(".")[-1]) <= 3 for r in rows)


def test_alarms_csv(report: DailyReport, tmp_path: Path) -> None:
    path = tmp_path / "alarms.csv"
    write_alarms_csv(report, path)
    rows = list(csv.DictReader(path.open(encoding="utf-8"), delimiter=";"))
    assert len(rows) == report.alarms_total
    assert list(rows[0]) == ["ts", "code", "device", "slot", "priority", "text"]
    assert rows[0]["ts"].endswith("+03:00")
    assert any(r["code"] == "P102.FTS" and r["text"] == "Насос Н2: нет подтверждения пуска" for r in rows)


def _cf_ranges(ws: Worksheet) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {}
    for rng in ws.conditional_formatting:
        result[str(rng.sqref)] = [rule.type or "" for rule in rng.rules]
    return result


def test_xlsx_structure_and_styles(report: DailyReport, tmp_path: Path) -> None:
    paths = write_outputs(report, tmp_path, ["xlsx"])
    assert [p.name for p in paths] == ["kns1_daily_2026-10-15.xlsx"]
    wb = load_workbook(paths[0])
    assert wb.sheetnames == [SHEET_HOURLY, SHEET_SUMMARY, SHEET_ALARMS, SHEET_CHARTS]

    ws = wb[SHEET_HOURLY]
    assert "суточный отчёт за 15.10.2026" in str(ws["A1"].value)
    header = [c.value for c in ws[HEADER_ROW]]
    assert header[:5] == ["Час", "Уровень мин., м", "Уровень ср., м", "Уровень макс., м", "Откачано, м³"]
    assert "Н1: работа, ч" in header and "Полнота данных" in header
    head = ws.cell(row=HEADER_ROW, column=2)
    assert head.font.bold and head.fill.fgColor.rgb.endswith("37474F")
    assert ws.freeze_panes == f"B{FIRST_DATA_ROW}"

    first = FIRST_DATA_ROW
    assert ws.cell(row=first, column=1).value == "00:00–01:00"
    assert ws.cell(row=first, column=2).number_format == "0.00"
    assert ws.cell(row=first, column=5).number_format == "0.0"
    assert ws.cell(row=first, column=6).number_format == "0.00"
    assert ws.cell(row=first, column=7).number_format == "0"
    assert ws.cell(row=first, column=16).number_format == "0%"
    assert ws.cell(row=first + 3, column=16).value == pytest.approx(report.hours[3].coverage)

    total = first + 24
    assert ws.cell(row=total, column=1).value == "Сутки"
    assert ws.cell(row=total, column=5).value == pytest.approx(report.volume_m3)
    assert ws.cell(row=total, column=12).value == report.alarms_total

    cf = _cf_ranges(ws)
    last = first + 23
    assert cf[f"L{first}:L{last}"] == ["cellIs"]  # тревог всего > 0
    assert cf[f"M{first}:M{last}"] == ["cellIs"]  # High > 0
    assert cf[f"D{first}:D{last}"] == ["cellIs", "cellIs"]  # уровень макс. ≥ HH, ≥ H
    assert cf[f"B{first}:B{last}"] == ["expression", "expression"]  # уровень мин. ≤ LL, ≤ L
    assert cf[f"P{first}:P{last}"] == ["cellIs"]  # полнота < 95 %
    assert cf[f"E{first}:E{last}"] == ["dataBar"]
    rules = {rule.formula[0] for rng in ws.conditional_formatting for rule in rng.rules if rule.formula}
    assert {"4.0", "3.5", "0.95", "0"} <= rules


def test_xlsx_summary_alarms_and_charts(report: DailyReport, tmp_path: Path) -> None:
    path = write_outputs(report, tmp_path, ["xlsx"])[0]
    wb = load_workbook(path)

    summary = wb[SHEET_SUMMARY]
    values = {row[0].value: row[1].value for row in summary.iter_rows(min_row=3, max_row=20) if row[0].value}
    assert values["Откачано, м³"] == pytest.approx(report.volume_m3)
    assert values["Тревог всего"] == report.alarms_total
    pump_rows = [r for r in summary.iter_rows(values_only=True) if r[0] and str(r[0]).startswith("Насос Н")]
    assert [r[0] for r in pump_rows] == ["Насос Н1 (P101)", "Насос Н2 (P102)", "Насос Н3 (P103)"]
    assert pump_rows[0][5] == "счётчики ПЛК / счётчики ПЛК"

    alarms_ws = wb[SHEET_ALARMS]
    assert alarms_ws.max_row == report.alarms_total + 1
    assert alarms_ws.auto_filter.ref == f"A1:E{report.alarms_total + 1}"
    assert alarms_ws.cell(row=2, column=1).number_format == "dd.mm.yyyy hh:mm:ss"

    with zipfile.ZipFile(path) as archive:
        charts = [n for n in archive.namelist() if n.startswith("xl/charts/chart")]
    assert len(charts) == 4


def test_xlsx_without_alarms_and_partial_period(
    cfg: ReportConfig, tags: dict[str, TagInfo], alarms: list[AlarmDef], tmp_path: Path
) -> None:
    from kns_report.samples import load_jsonl
    from kns_report.timeutil import ceil_hour, floor_hour

    series = load_jsonl(FIXTURES / "opcua_window.jsonl")
    span = series.time_range()
    assert span is not None
    quiet = build_report(
        series, cfg.station, tags, alarms, floor_hour(span[0], MSK), ceil_hour(span[1], MSK), "opcua", GENERATED_AT
    )
    paths = write_outputs(quiet, tmp_path, ["csv", "xlsx"])
    assert [p.name for p in paths] == [
        "kns1_report_2026-10-15_1400-1500.csv",
        "kns1_report_2026-10-15_1400-1500_alarms.csv",
        "kns1_report_2026-10-15_1400-1500.xlsx",
    ]
    wb = load_workbook(paths[-1])
    assert "за период 15.10.2026 14:00" in str(wb[SHEET_HOURLY]["A1"].value)
    assert wb[SHEET_HOURLY].cell(row=FIRST_DATA_ROW + 1, column=1).value == "Итого"
    assert quiet.alarms_total == 1  # P102.FTS по фронту слова P102.AlmActive
    assert wb[SHEET_ALARMS]["B2"].value == "P102.FTS"


def test_xlsx_alarms_sheet_when_quiet(report: DailyReport, tmp_path: Path) -> None:
    report.events.clear()
    path = write_outputs(report, tmp_path, ["xlsx"])[0]
    assert load_workbook(path)[SHEET_ALARMS]["A2"].value == "Срабатываний тревог за период нет"
