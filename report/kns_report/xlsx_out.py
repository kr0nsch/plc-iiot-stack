"""XLSX отчёта (openpyxl): почасовая таблица, сводка, тревоги и графики.

Оформление в духе ISA-101: нейтральная таблица, цвет — только для отклонений (уставки уровня,
тревоги, неполные данные). Значения записываются числами (не формулами), чтобы файл читался
любой программой без пересчёта.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from kns_report.aggregate import PRIORITY_BUCKETS, DailyReport
from kns_report.table import Column, hourly_columns

SHEET_HOURLY = "Отчёт"
SHEET_SUMMARY = "Сводка"
SHEET_ALARMS = "Тревоги"
SHEET_CHARTS = "Графики"

HEADER_ROW = 4
FIRST_DATA_ROW = HEADER_ROW + 1

NUMBER_FORMATS = {
    "level": "0.00",
    "volume": "0.0",
    "hours": "0.00",
    "count": "0",
    "alarm": "0",
    "alarm_high": "0",
    "pct": "0%",
}

_HEADER_FILL = PatternFill("solid", fgColor="37474F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TOTAL_FILL = PatternFill("solid", fgColor="ECEFF1")
_THIN = Side(style="thin", color="B0BEC5")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_TITLE_FONT = Font(bold=True, size=14)
_NOTE_FONT = Font(italic=True, color="607D8B")
# Классическая палитра условного форматирования Excel
_BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
_BAD_FONT = Font(color="9C0006")
_ALARM_FILL = PatternFill("solid", fgColor="C00000")
_ALARM_FONT = Font(bold=True, color="FFFFFF")
_WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
_WARN_FONT = Font(color="9C5700")
_PRIORITY_COLORS = {"High": "C00000", "Medium": "ED7D31", "Low": "FFC000"}


def _date_label(report: DailyReport) -> str:
    start = report.period_start
    end = report.period_end
    if (end - start).total_seconds() == 86400 and start.hour == 0:
        return f"за {start:%d.%m.%Y}"
    return f"за период {start:%d.%m.%Y %H:%M} – {end:%d.%m.%Y %H:%M}"


def _tz_label(report: DailyReport) -> str:
    offset = report.period_start.utcoffset()
    if offset is None:
        return ""
    minutes = int(offset.total_seconds() // 60)
    sign = "+" if minutes >= 0 else "-"
    return f"UTC{sign}{abs(minutes) // 60:02d}:{abs(minutes) % 60:02d}"


def _style_header(ws: Worksheet, row: int, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_hourly(ws: Worksheet, report: DailyReport, columns: list[Column]) -> int:
    """Таблица по часам; возвращает номер строки итогов."""
    last_col = get_column_letter(len(columns))
    ws["A1"] = f"{report.station_name} — суточный отчёт {_date_label(report)}"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells(f"A1:{last_col}1")
    generated = report.generated_at.astimezone(report.period_start.tzinfo)
    ws["A2"] = (
        f"Источник: {report.source}. Время местное ({_tz_label(report)}). Сформирован {generated:%d.%m.%Y %H:%M}."
    )
    ws["A2"].font = _NOTE_FONT
    ws.merge_cells(f"A2:{last_col}2")

    for index, column in enumerate(columns, 1):
        ws.cell(row=HEADER_ROW, column=index, value=column.title)
    _style_header(ws, HEADER_ROW, len(columns))
    ws.row_dimensions[HEADER_ROW].height = 32

    for offset, hour in enumerate(report.hours):
        row_number = FIRST_DATA_ROW + offset
        for index, column in enumerate(columns, 1):
            if column.kind == "time":
                value: float | int | str | None = f"{hour.start:%H:%M}–{hour.end:%H:%M}"
            elif column.kind == "pct":
                value = hour.coverage
            else:
                value = column.get(hour)
            cell = ws.cell(row=row_number, column=index, value=value)
            cell.border = _BORDER
            if column.kind in NUMBER_FORMATS:
                cell.number_format = NUMBER_FORMATS[column.kind]
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="center")

    total_row = FIRST_DATA_ROW + len(report.hours)
    totals = _totals(report, columns)
    for index, column in enumerate(columns, 1):
        cell = ws.cell(row=total_row, column=index, value=totals.get(column.key))
        cell.font = Font(bold=True)
        cell.fill = _TOTAL_FILL
        cell.border = _BORDER
        if column.kind in NUMBER_FORMATS:
            cell.number_format = NUMBER_FORMATS[column.kind]
    ws.cell(row=total_row, column=1).alignment = Alignment(horizontal="center")

    widths = {"time": 13, "level": 11, "volume": 11, "hours": 10, "count": 9, "alarm": 9, "alarm_high": 9, "pct": 10}
    for index, column in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(column.kind, 10)
    ws.freeze_panes = f"B{FIRST_DATA_ROW}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    if ws.sheet_properties.pageSetUpPr is not None:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    _conditional(ws, report, columns, total_row - 1)
    return total_row


def _totals(report: DailyReport, columns: list[Column]) -> dict[str, float | int | str | None]:
    totals: dict[str, float | int | str | None] = {
        "hour_start": "Сутки" if len(report.hours) == 24 else "Итого",
        "level_min_m": report.level_min_m,
        "level_avg_m": report.level_avg_m,
        "level_max_m": report.level_max_m,
        "volume_m3": report.volume_m3,
        "alarms_total": report.alarms_total,
        "alarms_high": report.alarms_by_priority.get("High", 0),
        "alarms_medium": report.alarms_by_priority.get("Medium", 0),
        "alarms_low": report.alarms_by_priority.get("Low", 0),
        "data_coverage_pct": report.coverage,
    }
    for pump in report.pumps:
        key = pump.device.lower()
        totals[f"{key}_run_h"] = pump.run_h
        totals[f"{key}_starts"] = pump.starts
    return {c.key: totals.get(c.key) for c in columns}


def _conditional(ws: Worksheet, report: DailyReport, columns: list[Column], last_row: int) -> None:
    def col_range(key: str) -> str | None:
        for index, column in enumerate(columns, 1):
            if column.key == key:
                letter = get_column_letter(index)
                return f"{letter}{FIRST_DATA_ROW}:{letter}{last_row}"
        return None

    limits = report.level_limits
    high_range = col_range("level_max_m")
    if high_range:
        if "HH" in limits:
            ws.conditional_formatting.add(
                high_range,
                CellIsRule(
                    operator="greaterThanOrEqual",
                    formula=[repr(limits["HH"])],
                    fill=_ALARM_FILL,
                    font=_ALARM_FONT,
                    stopIfTrue=True,
                ),
            )
        if "H" in limits:
            ws.conditional_formatting.add(
                high_range,
                CellIsRule(
                    operator="greaterThanOrEqual", formula=[repr(limits["H"])], fill=_WARN_FILL, font=_WARN_FONT
                ),
            )
    low_range = col_range("level_min_m")
    if low_range:
        if "LL" in limits:
            ws.conditional_formatting.add(
                low_range,
                FormulaRule(
                    formula=[f"AND(ISNUMBER({low_range.split(':')[0]}),{low_range.split(':')[0]}<={limits['LL']!r})"],
                    fill=_ALARM_FILL,
                    font=_ALARM_FONT,
                    stopIfTrue=True,
                ),
            )
        if "L" in limits:
            ws.conditional_formatting.add(
                low_range,
                FormulaRule(
                    formula=[f"AND(ISNUMBER({low_range.split(':')[0]}),{low_range.split(':')[0]}<={limits['L']!r})"],
                    fill=_WARN_FILL,
                    font=_WARN_FONT,
                ),
            )
    alarm_range = col_range("alarms_total")
    if alarm_range:
        ws.conditional_formatting.add(
            alarm_range, CellIsRule(operator="greaterThan", formula=["0"], fill=_BAD_FILL, font=_BAD_FONT)
        )
    high_alarm_range = col_range("alarms_high")
    if high_alarm_range:
        ws.conditional_formatting.add(
            high_alarm_range, CellIsRule(operator="greaterThan", formula=["0"], fill=_ALARM_FILL, font=_ALARM_FONT)
        )
    coverage_range = col_range("data_coverage_pct")
    if coverage_range:
        ws.conditional_formatting.add(
            coverage_range, CellIsRule(operator="lessThan", formula=["0.95"], fill=_WARN_FILL, font=_WARN_FONT)
        )
    volume_range = col_range("volume_m3")
    if volume_range:
        ws.conditional_formatting.add(
            volume_range,
            DataBarRule(start_type="num", start_value=0, end_type="max", color="90A4AE", showValue=True),
        )


def _write_summary(ws: Worksheet, report: DailyReport) -> None:
    ws["A1"] = f"{report.station_name} — сводка {_date_label(report)}"
    ws["A1"].font = _TITLE_FONT
    rows: list[tuple[str, float | int | str | None, str]] = [
        ("Объект", f"{report.station_name} ({report.station_code})", ""),
        (
            "Период",
            f"{report.period_start:%d.%m.%Y %H:%M} – {report.period_end:%d.%m.%Y %H:%M} ({_tz_label(report)})",
            "",
        ),
        ("Источник данных", report.source, ""),
        ("Уровень мин., м", report.level_min_m, "0.00"),
        ("Уровень ср., м", report.level_avg_m, "0.00"),
        ("Уровень макс., м", report.level_max_m, "0.00"),
        ("Откачано, м³", report.volume_m3, "0.0"),
        ("Тревог всего", report.alarms_total, "0"),
    ]
    rows += [(f"  из них {p}", report.alarms_by_priority.get(p, 0), "0") for p in PRIORITY_BUCKETS]
    rows.append(("Полнота данных (уровень и расход)", report.coverage, "0%"))
    for offset, (label, value, fmt) in enumerate(rows, 3):
        ws.cell(row=offset, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=offset, column=2, value=value)
        if fmt:
            cell.number_format = fmt
        cell.alignment = Alignment(horizontal="left")

    start = 3 + len(rows) + 1
    headers = [
        "Насос",
        "Работа за период, ч",
        "Пуски за период",
        "Наработка всего на конец, ч",
        "Пусков всего на конец",
        "Метод (работа / пуски)",
    ]
    for index, title in enumerate(headers, 1):
        ws.cell(row=start, column=index, value=title)
    _style_header(ws, start, len(headers))
    methods = {"counters": "счётчики ПЛК", "state": "по Running", "none": "нет данных"}
    for offset, pump in enumerate(report.pumps, start + 1):
        values: list[tuple[float | int | str | None, str]] = [
            (f"{pump.title} ({pump.device})", ""),
            (pump.run_h, "0.00"),
            (pump.starts, "0"),
            (pump.run_total_h, "0.0"),
            (pump.starts_total, "0"),
            (f"{methods[pump.run_method]} / {methods[pump.starts_method]}", ""),
        ]
        for index, (value, fmt) in enumerate(values, 1):
            cell = ws.cell(row=offset, column=index, value=value)
            cell.border = _BORDER
            if fmt:
                cell.number_format = fmt
    note_row = start + len(report.pumps) + 2
    if report.warnings:
        ws.cell(row=note_row, column=1, value="Замечания к данным").font = Font(bold=True, color="9C5700")
        for offset, text in enumerate(report.warnings, note_row + 1):
            ws.cell(row=offset, column=1, value=f"• {text}").font = _WARN_FONT
    for letter, width in zip("ABCDEF", (36, 22, 16, 26, 22, 30), strict=True):
        ws.column_dimensions[letter].width = width


def _write_alarms(ws: Worksheet, report: DailyReport) -> None:
    headers = ["Время", "Код", "Устройство", "Приоритет", "Текст"]
    for index, title in enumerate(headers, 1):
        ws.cell(row=1, column=index, value=title)
    _style_header(ws, 1, len(headers))
    tz = report.period_start.tzinfo
    for offset, event in enumerate(report.events, 2):
        local = event.ts.astimezone(tz).replace(tzinfo=None)
        values: list[datetime | str] = [local, event.code, event.device, event.priority, event.text]
        for index, value in enumerate(values, 1):
            cell = ws.cell(row=offset, column=index, value=value)
            cell.border = _BORDER
        ws.cell(row=offset, column=1).number_format = "dd.mm.yyyy hh:mm:ss"
    last = max(2, len(report.events) + 1)
    if report.events:
        ws.auto_filter.ref = f"A1:E{last}"
        for priority, color in _PRIORITY_COLORS.items():
            ws.conditional_formatting.add(
                f"D2:D{last}",
                FormulaRule(formula=[f'$D2="{priority}"'], font=Font(bold=True, color=color)),
            )
    else:
        ws.cell(row=2, column=1, value="Срабатываний тревог за период нет").font = _NOTE_FONT
    ws.freeze_panes = "A2"
    for letter, width in zip("ABCDE", (20, 14, 12, 11, 70), strict=True):
        ws.column_dimensions[letter].width = width


def _write_charts(ws: Worksheet, data: Worksheet, report: DailyReport, columns: list[Column]) -> None:
    n = len(report.hours)
    if n == 0:
        ws["A1"] = "Нет данных для графиков"
        return
    first, last = FIRST_DATA_ROW, FIRST_DATA_ROW + n - 1
    index_of = {c.key: i for i, c in enumerate(columns, 1)}
    categories = Reference(data, min_col=1, min_row=first, max_row=last)

    def ref(key: str) -> Reference:
        col = index_of[key]
        return Reference(data, min_col=col, min_row=HEADER_ROW, max_row=last)

    level = LineChart()
    level.title = "Уровень в резервуаре LT101, м"
    level.y_axis.title = "м"
    level.x_axis.title = "Час"
    for key, color, width in (
        ("level_max_m", "90A4AE", 15000),
        ("level_avg_m", "263238", 28000),
        ("level_min_m", "90A4AE", 15000),
    ):
        level.add_data(ref(key), titles_from_data=True)
        series = level.series[-1]
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = width
        series.smooth = False
    level.set_categories(categories)

    volume = BarChart()
    volume.type = "col"
    volume.title = "Откачано по часам, м³"
    volume.y_axis.title = "м³"
    volume.add_data(ref("volume_m3"), titles_from_data=True)
    volume.series[0].graphicalProperties.solidFill = "546E7A"
    volume.set_categories(categories)
    volume.legend = None

    pumps = BarChart()
    pumps.type = "col"
    pumps.grouping = "stacked"
    pumps.overlap = 100
    pumps.title = "Работа насосов, ч"
    pumps.y_axis.title = "ч"
    for pump, color in zip(report.pumps, ("37474F", "78909C", "B0BEC5", "CFD8DC"), strict=False):
        pumps.add_data(ref(f"{pump.device.lower()}_run_h"), titles_from_data=True)
        pumps.series[-1].graphicalProperties.solidFill = color
    pumps.set_categories(categories)

    alarms = BarChart()
    alarms.type = "col"
    alarms.grouping = "stacked"
    alarms.overlap = 100
    alarms.title = "Срабатывания тревог по приоритетам"
    for key, color in (("alarms_high", "C00000"), ("alarms_medium", "ED7D31"), ("alarms_low", "FFC000")):
        alarms.add_data(ref(key), titles_from_data=True)
        alarms.series[-1].graphicalProperties.solidFill = color
    alarms.set_categories(categories)

    for chart, anchor in ((level, "A1"), (volume, "K1"), (pumps, "A22"), (alarms, "K22")):
        chart.width = 17
        chart.height = 10
        # openpyxl 3.1 по умолчанию помечает оси удалёнными — Excel их тогда не рисует
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        ws.add_chart(chart, anchor)


def write_xlsx(report: DailyReport, path: Path) -> None:
    columns = hourly_columns(report)
    wb = Workbook()
    hourly = wb.active
    assert hourly is not None
    hourly.title = SHEET_HOURLY
    _write_hourly(hourly, report, columns)
    _write_summary(wb.create_sheet(SHEET_SUMMARY), report)
    _write_alarms(wb.create_sheet(SHEET_ALARMS), report)
    _write_charts(wb.create_sheet(SHEET_CHARTS), hourly, report, columns)
    wb.properties.title = f"{report.station_name}: суточный отчёт"
    wb.properties.creator = "kns-report"
    wb.save(path)
